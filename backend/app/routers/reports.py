"""Hisobotlar: tushum, foyda, xizmatlar, ustalar, ombor + Excel eksport."""

from datetime import date, datetime, time, timedelta
from io import BytesIO

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import distinct, func, select
from sqlalchemy.orm import Session

from app.core.deps import DB, StaffUser
from app.models import (
    Appointment,
    AppointmentStatus,
    Customer,
    Expense,
    Order,
    OrderPart,
    OrderService,
    OrderStatus,
    Part,
    Payment,
    User,
)
from app.services.orders import OPEN_STATUSES
from app.utils import now

router = APIRouter(prefix="/reports", tags=["reports"])

METHOD_LABELS = {"cash": "Naqd", "card": "Karta", "transfer": "O'tkazma", "click": "Click", "payme": "Payme"}


def _f(v) -> float:
    return float(v or 0)


def _range(date_from: date | None, date_to: date | None) -> tuple[datetime, datetime, date, date]:
    today = now().date()
    d_from = date_from or today.replace(day=1)
    d_to = date_to or today
    if d_from > d_to:
        raise HTTPException(400, "Boshlanish sanasi tugash sanasidan katta")
    if (d_to - d_from).days > 366 * 3:
        raise HTTPException(400, "Davr juda katta")
    return datetime.combine(d_from, time.min), datetime.combine(d_to, time.max), d_from, d_to


def _closed_in(dt_from, dt_to):
    return (Order.status == OrderStatus.closed, Order.closed_at >= dt_from, Order.closed_at <= dt_to)


@router.get("/dashboard")
def dashboard(db: DB, _: StaffUser):
    current = now()
    today_start = datetime.combine(current.date(), time.min)
    month_start = today_start.replace(day=1)

    def paid_since(start):
        return _f(db.scalar(select(func.sum(Payment.amount)).where(Payment.created_at >= start)))

    status_counts = dict(
        db.execute(
            select(Order.status, func.count(Order.id)).where(Order.status.in_(OPEN_STATUSES)).group_by(Order.status)
        ).all()
    )
    appointments_today = db.scalar(
        select(func.count(Appointment.id)).where(
            Appointment.scheduled_at >= today_start,
            Appointment.scheduled_at < today_start + timedelta(days=1),
            Appointment.status.not_in((AppointmentStatus.cancelled, AppointmentStatus.no_show)),
        )
    )
    debt = db.scalar(
        select(func.sum(Order.total - Order.paid)).where(Order.status != OrderStatus.cancelled, Order.total > Order.paid)
    )
    return {
        "revenue_today": paid_since(today_start),
        "revenue_month": paid_since(month_start),
        "orders_today": db.scalar(select(func.count(Order.id)).where(Order.created_at >= today_start)),
        "open_orders": {str(k): v for k, v in status_counts.items()},
        "appointments_today": appointments_today,
        "low_stock": db.scalar(
            select(func.count(Part.id)).where(Part.is_active.is_(True), Part.quantity <= Part.min_quantity)
        ),
        "total_debt": _f(debt),
        "customers_total": db.scalar(select(func.count(Customer.id))),
    }


def summary_data(db: Session, dt_from: datetime, dt_to: datetime) -> dict:
    closed = _closed_in(dt_from, dt_to)
    row = db.execute(
        select(
            func.count(Order.id),
            func.coalesce(func.sum(Order.total), 0),
            func.coalesce(func.sum(Order.services_total), 0),
            func.coalesce(func.sum(Order.parts_total), 0),
            func.coalesce(func.sum(Order.discount), 0),
        ).where(*closed)
    ).one()
    orders_count, revenue, services_rev, parts_rev, discount = row
    parts_cost = db.scalar(
        select(func.coalesce(func.sum(OrderPart.cost * OrderPart.quantity), 0)).join(Order).where(*closed)
    )
    payments = db.execute(
        select(Payment.method, func.sum(Payment.amount))
        .where(Payment.created_at >= dt_from, Payment.created_at <= dt_to)
        .group_by(Payment.method)
    ).all()
    expenses = db.execute(
        select(Expense.category, func.sum(Expense.amount))
        .where(Expense.spent_on >= dt_from.date(), Expense.spent_on <= dt_to.date())
        .group_by(Expense.category)
    ).all()
    commissions = sum(m["commission"] for m in mechanics_data(db, dt_from, dt_to))
    expenses_total = sum(_f(v) for _, v in expenses)
    gross_profit = _f(revenue) - _f(parts_cost)
    return {
        "orders_closed": orders_count,
        "revenue": _f(revenue),
        "services_revenue": _f(services_rev),
        "parts_revenue": _f(parts_rev),
        "parts_cost": _f(parts_cost),
        "discount": _f(discount),
        "avg_check": _f(revenue) / orders_count if orders_count else 0,
        "gross_profit": gross_profit,
        "mechanic_commissions": commissions,
        "expenses_total": expenses_total,
        "net_profit": gross_profit - commissions - expenses_total,
        "cash_in": sum(_f(v) for _, v in payments),
        "payments_by_method": [{"method": str(m), "label": METHOD_LABELS.get(str(m), str(m)), "amount": _f(v)} for m, v in payments],
        "expenses_by_category": [{"category": c, "amount": _f(v)} for c, v in expenses],
        "new_customers": db.scalar(
            select(func.count(Customer.id)).where(Customer.created_at >= dt_from, Customer.created_at <= dt_to)
        ),
    }


def daily_data(db: Session, dt_from: datetime, dt_to: datetime) -> list[dict]:
    pay_day = func.date(Payment.created_at)
    payments = {
        str(d)[:10]: _f(v)
        for d, v in db.execute(
            select(pay_day, func.sum(Payment.amount))
            .where(Payment.created_at >= dt_from, Payment.created_at <= dt_to)
            .group_by(pay_day)
        ).all()
    }
    close_day = func.date(Order.closed_at)
    closed = {
        str(d)[:10]: (c, _f(v))
        for d, c, v in db.execute(
            select(close_day, func.count(Order.id), func.sum(Order.total))
            .where(*_closed_in(dt_from, dt_to))
            .group_by(close_day)
        ).all()
    }
    result = []
    d = dt_from.date()
    while d <= dt_to.date():
        key = d.isoformat()
        count, total = closed.get(key, (0, 0.0))
        result.append({"date": key, "cash_in": payments.get(key, 0.0), "orders_closed": count, "revenue": total})
        d += timedelta(days=1)
    return result


def services_data(db: Session, dt_from: datetime, dt_to: datetime) -> list[dict]:
    rows = db.execute(
        select(
            OrderService.name,
            func.sum(OrderService.quantity),
            func.sum(OrderService.price * OrderService.quantity),
        )
        .join(Order)
        .where(*_closed_in(dt_from, dt_to))
        .group_by(OrderService.name)
        .order_by(func.sum(OrderService.price * OrderService.quantity).desc())
    ).all()
    return [{"name": n, "count": _f(c), "revenue": _f(r)} for n, c, r in rows]


def mechanics_data(db: Session, dt_from: datetime, dt_to: datetime) -> list[dict]:
    rows = db.execute(
        select(
            User.id,
            User.full_name,
            User.commission_percent,
            func.count(distinct(Order.id)),
            func.sum(OrderService.price * OrderService.quantity),
        )
        .select_from(OrderService)
        .join(Order, OrderService.order_id == Order.id)
        .join(User, OrderService.mechanic_id == User.id)
        .where(*_closed_in(dt_from, dt_to))
        .group_by(User.id, User.full_name, User.commission_percent)
        .order_by(func.sum(OrderService.price * OrderService.quantity).desc())
    ).all()
    return [
        {
            "id": uid,
            "name": name,
            "orders": cnt,
            "revenue": _f(rev),
            "percent": _f(pct),
            "commission": round(_f(rev) * _f(pct) / 100, 2),
        }
        for uid, name, pct, cnt, rev in rows
    ]


def parts_data(db: Session, dt_from: datetime, dt_to: datetime) -> list[dict]:
    rows = db.execute(
        select(
            OrderPart.name,
            func.sum(OrderPart.quantity),
            func.sum(OrderPart.price * OrderPart.quantity),
            func.sum(OrderPart.cost * OrderPart.quantity),
        )
        .join(Order)
        .where(*_closed_in(dt_from, dt_to))
        .group_by(OrderPart.name)
        .order_by(func.sum(OrderPart.price * OrderPart.quantity).desc())
    ).all()
    return [
        {"name": n, "quantity": _f(q), "revenue": _f(r), "cost": _f(c), "profit": _f(r) - _f(c)} for n, q, r, c in rows
    ]


@router.get("/summary")
def summary(db: DB, _: StaffUser, date_from: date | None = None, date_to: date | None = None):
    dt_from, dt_to, d_from, d_to = _range(date_from, date_to)
    return {"date_from": d_from, "date_to": d_to, **summary_data(db, dt_from, dt_to)}


@router.get("/daily")
def daily(db: DB, _: StaffUser, date_from: date | None = None, date_to: date | None = None):
    dt_from, dt_to, *_ = _range(date_from, date_to)
    return daily_data(db, dt_from, dt_to)


@router.get("/services")
def services(db: DB, _: StaffUser, date_from: date | None = None, date_to: date | None = None):
    dt_from, dt_to, *_ = _range(date_from, date_to)
    return services_data(db, dt_from, dt_to)


@router.get("/mechanics")
def mechanics(db: DB, _: StaffUser, date_from: date | None = None, date_to: date | None = None):
    dt_from, dt_to, *_ = _range(date_from, date_to)
    return mechanics_data(db, dt_from, dt_to)


@router.get("/parts")
def parts(db: DB, _: StaffUser, date_from: date | None = None, date_to: date | None = None):
    dt_from, dt_to, *_ = _range(date_from, date_to)
    return parts_data(db, dt_from, dt_to)


# ---------------- Excel ----------------

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append(row)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None] + [8])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(width + 2, 50)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0"
    return ws


@router.get("/export")
def export_xlsx(db: DB, _: StaffUser, date_from: date | None = None, date_to: date | None = None):
    dt_from, dt_to, d_from, d_to = _range(date_from, date_to)
    s = summary_data(db, dt_from, dt_to)
    wb = Workbook()
    _sheet(
        wb,
        "Umumiy",
        ["Ko'rsatkich", "Qiymat"],
        [
            ["Davr", f"{d_from:%d.%m.%Y} — {d_to:%d.%m.%Y}"],
            ["Yopilgan buyurtmalar", s["orders_closed"]],
            ["Tushum (buyurtmalar)", s["revenue"]],
            ["  shundan xizmatlar", s["services_revenue"]],
            ["  shundan ehtiyot qismlar", s["parts_revenue"]],
            ["Chegirmalar", s["discount"]],
            ["Ehtiyot qismlar tannarxi", s["parts_cost"]],
            ["Yalpi foyda", s["gross_profit"]],
            ["Ustalar ish haqi", s["mechanic_commissions"]],
            ["Xarajatlar", s["expenses_total"]],
            ["Sof foyda", s["net_profit"]],
            ["O'rtacha chek", s["avg_check"]],
            ["Kassaga tushgan pul", s["cash_in"]],
            *[[f"  {p['label']}", p["amount"]] for p in s["payments_by_method"]],
            ["Yangi mijozlar", s["new_customers"]],
        ],
        first=True,
    )
    _sheet(
        wb,
        "Kunlik",
        ["Sana", "Kassaga tushgan", "Yopilgan buyurtmalar", "Tushum"],
        [[d["date"], d["cash_in"], d["orders_closed"], d["revenue"]] for d in daily_data(db, dt_from, dt_to)],
    )
    _sheet(
        wb,
        "Xizmatlar",
        ["Xizmat", "Soni", "Summa"],
        [[r["name"], r["count"], r["revenue"]] for r in services_data(db, dt_from, dt_to)],
    )
    _sheet(
        wb,
        "Ustalar",
        ["Usta", "Buyurtmalar", "Bajargan ishlar summasi", "Foiz", "Ish haqi"],
        [[r["name"], r["orders"], r["revenue"], r["percent"], r["commission"]] for r in mechanics_data(db, dt_from, dt_to)],
    )
    _sheet(
        wb,
        "Ehtiyot qismlar",
        ["Nomi", "Miqdori", "Sotuv summasi", "Tannarx", "Foyda"],
        [[r["name"], r["quantity"], r["revenue"], r["cost"], r["profit"]] for r in parts_data(db, dt_from, dt_to)],
    )
    _sheet(
        wb,
        "Xarajatlar",
        ["Kategoriya", "Summa"],
        [[e["category"], e["amount"]] for e in s["expenses_by_category"]],
    )
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"hisobot_{d_from:%Y%m%d}_{d_to:%Y%m%d}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

